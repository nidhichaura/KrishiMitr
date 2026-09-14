from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("output/KrishiMitr_Hackathon_Product_Documentation.docx")
BLUE = "0B4F8A"
LIGHT_BLUE = "EAF5FF"
PALE = "F6FAFD"
GREY = "D9E2EA"
TEXT = "1A2B3B"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:fill"), fill)
    tc_pr.append(el)


def borders(cell, color=GREY):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = tc_borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, color=TEXT, size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Aptos"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    borders(cell)


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.autofit = False
    for i, header in enumerate(headers):
        cell = t.rows[0].cells[i]
        if widths:
            cell.width = Inches(widths[i])
        shade(cell, BLUE)
        set_cell_text(cell, header, bold=True, color="FFFFFF", size=8)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            if widths:
                cells[i].width = Inches(widths[i])
            if ri % 2:
                shade(cells[i], PALE)
            set_cell_text(cells[i], value, size=8)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return t


def base_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.1
    for name, size in [("Title", 28), ("Heading 1", 18), ("Heading 2", 12)]:
        s = doc.styles[name]
        s.font.name = "Aptos Display" if name != "Normal" else "Aptos"
        s._element.rPr.rFonts.set(qn("w:ascii"), s.font.name)
        s._element.rPr.rFonts.set(qn("w:hAnsi"), s.font.name)
        s.font.size = Pt(size)
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.font.bold = name != "Normal"
        s.paragraph_format.space_before = Pt(10)
        s.paragraph_format.space_after = Pt(7)


def add_header_footer(doc):
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("KRISHIMITR  |  HACKATHON PRODUCT DOCUMENTATION")
    r.font.name = "Aptos"
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(BLUE)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("KrishiMitr  |  Version 1.0  |  Hackathon Submission")
    r.font.name = "Aptos"
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string("667788")


def page(doc, number, title, subtitle):
    doc.add_heading(f"{number:02d}  {title}", level=1)
    p = doc.add_paragraph(subtitle)
    p.style = doc.styles["Normal"]
    for r in p.runs:
        r.italic = True
        r.font.color.rgb = RGBColor.from_string("58708A")


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def main():
    OUT.parent.mkdir(exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)
    base_styles(doc)
    add_header_footer(doc)

    # 1 Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(66)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("HACKATHON")
    r.bold = True; r.font.size = Pt(15); r.font.color.rgb = RGBColor.from_string(BLUE)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("KrishiMitr Product Documentation")
    p = doc.add_paragraph("A complete product engineering and delivery record")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(13); p.runs[0].font.color.rgb = RGBColor.from_string("58708A")
    doc.add_paragraph()
    table(doc, ["PROJECT IDENTITY", "DETAIL"], [
        ["Product name", "KrishiMitr"], ["Team name", "KrishiMitr Team"],
        ["Project lead", "To be confirmed"], ["Team members", "To be confirmed"],
        ["Mentor", "To be confirmed"], ["Version", "1.0"],
        ["Submission date", "14 September 2026"],
    ], [2.0, 4.5])
    doc.add_heading("Document Control", level=2)
    table(doc, ["VERSION", "DATE", "OWNER", "CHANGE SUMMARY"], [["1.0", "14 Sep 2026", "KrishiMitr Team", "Initial hackathon submission"]], [0.8, 1.1, 1.8, 2.8])
    doc.add_page_break()

    # 2 Executive
    page(doc, 2, "Executive Summary", "The page a product leader should understand in under two minutes.")
    doc.add_heading("Product Brief", level=2)
    doc.add_paragraph("KrishiMitr is a multilingual, farmer-friendly web and WhatsApp assistant for small and marginal farmers in India. It helps farmers screen crop-leaf images for disease, retrieve verified mandi prices, receive location-aware crop guidance, and request follow-up support. The prototype reduces the friction of navigating fragmented agricultural information by delivering simple, local-language actions and clearly labelled government-source market data.")
    table(doc, ["DIMENSION", "DOCUMENTATION"], [
        ["Problem", "Farmers need timely crop health, market and advisory information but face language, trust and access barriers."],
        ["Target user", "Small and marginal farmers, especially mobile-first Hindi and regional-language users."],
        ["Core solution", "A single assistant combining leaf screening, mandi lookup, location-aware advice and WhatsApp access."],
        ["Primary value", "Faster, more understandable decisions using verified sources and transparent limitations."],
        ["Differentiator", "Farmer-first UX, live AGMARKNET data, multilingual interaction and consent-aware flows."],
    ], [1.55, 4.95])
    doc.add_heading("Vision and Success", level=2)
    doc.add_paragraph("KrishiMitr aims to make trusted agricultural decision support as easy to access as sending a WhatsApp message or pressing one clearly labelled button.")
    table(doc, ["KPI OR SUCCESS METRIC", "BASELINE", "TARGET", "MEASUREMENT METHOD"], [
        ["Verified price lookup", "Manual market search", "One guided request", "AGMARKNET response and UI result"],
        ["Advice completion", "Multi-step data entry", "Location-assisted flow", "Completed advisory requests"],
        ["Language accessibility", "English-heavy portals", "Hindi and English UI", "Language switch and TTS availability"],
    ], [1.8, 1.2, 1.4, 2.1])
    doc.add_page_break()

    # 3 Problem users
    page(doc, 3, "Problem Users and Opportunity", "Document the context before explaining the technology.")
    doc.add_heading("Problem Definition", level=2)
    doc.add_paragraph("A farmer making a same-day crop or selling decision often has to move between mandi portals, weather sources, crop-health advice and local contacts. Information may be unavailable in their preferred language, hard to validate, or unsuitable for a mobile screen. This creates delay and uncertainty at moments when price and crop-condition information matters most.")
    doc.add_heading("Users and Stakeholders", level=2)
    table(doc, ["PERSONA", "GOAL", "PAIN POINT", "SUCCESS LOOKS LIKE"], [
        ["Primary farmer", "Check crop health and market price quickly", "Complex forms, unclear sources, language barriers", "Clear answer, source and next action on mobile"],
        ["Buyer or trader", "Post or find crop requirements", "Disconnected local supply information", "Structured bid and nearby listing"],
        ["Agriculture advisor", "Support informed recommendations", "Incomplete farmer context", "Transparent field inputs and safety guardrails"],
        ["Pilot operator", "Improve the service responsibly", "Need consent-aware insights", "Aggregated signals without exposing personal data"],
    ], [1.25, 1.7, 1.9, 1.65])
    doc.add_heading("Existing Landscape", level=2)
    doc.add_paragraph("Government data portals provide valuable information but are not always designed as a guided farmer journey. Search, terminology, intermittent updates and lack of explanatory context can make verified information difficult to use at the field level.")
    table(doc, ["ALTERNATIVE", "STRENGTH", "LIMITATION", "KRISHIMITR ADVANTAGE"], [
        ["Government mandi portal", "Official price data", "Search-oriented and technical", "Guided crop and state selection with source shown"],
        ["Informal advice", "Accessible and familiar", "May be inconsistent or unverifiable", "Structured, traceable decision support"],
        ["Single-purpose crop apps", "Focused workflows", "Fragmented farmer journey", "Health, price and advice in one experience"],
    ], [1.4, 1.55, 1.7, 1.85])
    doc.add_page_break()

    # 4 Requirements
    page(doc, 4, "Product Requirements and Use Cases", "Translate the problem into a product that can be designed built and tested.")
    doc.add_heading("Core Use Cases", level=2)
    table(doc, ["ID", "ACTOR", "USE CASE", "TRIGGER OR INPUT", "EXPECTED OUTCOME"], [
        ["UC-01", "Farmer", "Leaf screening", "Clear crop leaf image", "Disease screening and safe next-step guidance"],
        ["UC-02", "Farmer", "Mandi price lookup", "Crop, state and optional mandi", "Verified price range, date and source"],
        ["UC-03", "Farmer", "Field advisory", "Location or field values", "Crop suitability and safe recommendation"],
        ["UC-04", "Buyer or seller", "Marketplace bid", "Crop, quantity, price and location", "Structured bid confirmation"],
        ["UC-05", "Farmer", "WhatsApp question", "Text or supported media", "Multilingual reply through Twilio"],
    ], [0.55, 0.85, 1.35, 1.8, 1.95])
    doc.add_heading("Functional Requirements", level=2)
    table(doc, ["ID", "REQUIREMENT", "PRIORITY", "ACCEPTANCE CRITERIA"], [
        ["FR-01", "Fetch verified mandi price", "High", "Returns crop, mandi, price range, date and source when official data exists"],
        ["FR-02", "Analyse leaf image", "High", "Accepts JPG, PNG or WEBP, validates quality and returns a disease result"],
        ["FR-03", "Support Hindi and English", "High", "Visible labels and key responses switch with the selected language"],
        ["FR-04", "Use location for weather", "Medium", "User permission fills weather values and shows source-aware locked state"],
        ["FR-05", "Capture optional consent", "Medium", "Core services work without consent; consent is explicit and visible"],
    ], [0.6, 2.5, 0.8, 2.6])
    doc.add_page_break()

    # 5 NFR / UX
    page(doc, 5, "Nonfunctional Requirements and User Journey", "Show how a real farmer experiences the product from entry to outcome.")
    table(doc, ["AREA", "REQUIREMENT", "TARGET OR STANDARD"], [
        ["Performance", "Fast feedback for common actions", "Visible loading state; external API timeouts fail safely"],
        ["Security", "Validate inputs and rate-limit public endpoints", "File types and sizes constrained; request rate limit applied"],
        ["Availability", "Protect core UI from external-source outages", "Clear verified-data unavailable message; no fabricated price"],
        ["Usability", "Mobile and low-literacy friendly", "Large controls, Hindi text, icons and read-aloud support"],
        ["Accessibility", "Keyboard and focus support", "Visible focus states and semantic labels"],
    ], [1.2, 2.5, 2.8])
    doc.add_heading("User Journey", level=2)
    table(doc, ["STAGE", "USER ACTION", "SYSTEM RESPONSE", "USER OUTCOME"], [
        ["1 Discover", "Opens KrishiMitr", "Shows two clear service entry points", "Understands available help"],
        ["2 Start", "Chooses mandi price", "Displays crop chips and state selector", "Can select without typing complex terms"],
        ["3 Execute", "Chooses Brinjal and Odisha", "Calls official market service", "Gets verified data or a transparent unavailable state"],
        ["4 Receive result", "Reads or listens to result", "Shows mandi, range, date and source", "Can act using a traceable price"],
        ["5 Continue", "Uses advisory or marketplace", "Retains guided, consent-aware flow", "Completes another relevant task"],
    ], [0.9, 1.65, 2.1, 1.85])
    doc.add_page_break()

    # 6 Architecture
    page(doc, 6, "System Architecture and Data Flow", "This is the engineering blueprint of the product.")
    doc.add_heading("Architecture Flow", level=2)
    table(doc, ["LAYER", "COMPONENT", "RESPONSIBILITY"], [
        ["Client", "Responsive web interface", "Farmer-facing Hindi and English UI, input validation, text-to-speech and location interaction"],
        ["Channel", "Twilio WhatsApp webhook", "Receives text and supported media from the WhatsApp channel"],
        ["API", "FastAPI routes", "Validates requests, rate-limits public endpoints and returns safe JSON or TwiML responses"],
        ["Business logic", "Market, advisory, vision and response services", "Normalises requests, applies rules and assembles explainable responses"],
        ["Data and AI", "AGMARKNET, Open-Meteo, PlantVillage model, SQLite", "Provides verified prices, weather, image screening and pilot records"],
    ], [1.0, 2.1, 3.5])
    doc.add_heading("End to End Workflow", level=2)
    doc.add_paragraph("Farmer request → client-side validation → FastAPI endpoint → service-specific validation → government or AI source → normalised result → language-aware response → optional consent-aware pilot storage. External sources are treated as dependencies: when a verified quote is unavailable, the product shows an explicit unavailable state instead of generating a price.")
    doc.add_heading("Trust Boundaries", level=2)
    bullet(doc, "Browser and WhatsApp inputs are untrusted until type, size and schema validation completes.")
    bullet(doc, "Official market and weather sources are external dependencies; their data is displayed with date and source context.")
    bullet(doc, "Consent-required pilot records are separated from public advisory and market lookup flows.")
    doc.add_page_break()

    # 7 Stack
    page(doc, 7, "Technology Stack and Engineering Design", "Explain not only what was used but why it was chosen.")
    table(doc, ["LAYER", "TECHNOLOGY", "WHY CHOSEN", "KEY RESPONSIBILITY"], [
        ["Frontend", "HTML CSS JavaScript", "Lightweight and responsive for mobile-first delivery", "Farmer UI, i18n, validation and interaction"],
        ["Backend", "Python FastAPI", "Typed API contracts and fast async web handling", "Routes, validation and service composition"],
        ["Database", "SQLite", "Simple pilot-ready local persistence", "Consent-aware profiles and marketplace bids"],
        ["AI and CV", "PyTorch EfficientNet and OpenCV", "Image classification and photo quality checks", "Leaf disease screening"],
        ["Integrations", "AGMARKNET Open-Meteo Twilio", "Authoritative market, weather and messaging access", "Prices, weather and WhatsApp channel"],
        ["Testing", "Pytest and syntax checks", "Repeatable confidence in public surface", "Route and behaviour checks"],
    ], [1.0, 1.6, 2.05, 1.95])
    doc.add_heading("Implementation Decisions", level=2)
    doc.add_paragraph("The market service prioritises AGMARKNET 2.0 and falls back to the legacy official data.gov.in feed only when configured. The response includes the official publication date, so the system does not represent older data as a new price. The advisory feature labels its output as decision support, not a fertilizer prescription.")
    doc.add_heading("Core Logic", level=2)
    doc.add_paragraph("The advisory service compares supplied weather and soil values with transparent crop profiles using a scaled nearest-profile distance. The market service normalises farmer crop terms, matches the official commodity spelling, filters by state and optional mandi, and returns the first valid market record from the most recent published report.")
    doc.add_page_break()

    # 8 Data privacy
    page(doc, 8, "Data APIs Security and Privacy", "Record how information is stored exchanged and protected.")
    doc.add_heading("Data Model", level=2)
    table(doc, ["ENTITY", "KEY FIELDS", "PURPOSE AND RETENTION"], [
        ["Farmer profile", "One-way identifier, state, district, language, consent", "Pilot interest and anonymous aggregate insights; consent-led"],
        ["Marketplace bid", "Trade type, crop, quantity, price, location, date", "Display of structured crop buying or selling intent"],
        ["Advisory request", "Crop, weather, NPK, pH, question", "Request-time decision support; not stored as a persistent farmer profile by the public route"],
        ["Leaf upload", "Uploaded image bytes", "Processed for screening during request; constrained by file type and size"],
        ["Operational logs", "Service events and errors", "Troubleshooting and availability monitoring"],
    ], [1.35, 2.25, 3.0])
    doc.add_heading("Security and Privacy Controls", level=2)
    bullet(doc, "Input schemas validate numeric ranges, required fields and upload media types before services run.")
    bullet(doc, "Rate limiting protects public endpoints from excessive request volume.")
    bullet(doc, "WhatsApp signature verification is configurable for production webhook authenticity.")
    bullet(doc, "Consent is optional for core services and explicitly captured before pilot data is used for anonymous improvement insights.")
    bullet(doc, "Secrets such as Twilio and data-provider keys are configured through environment variables rather than committed to source control.")
    doc.add_page_break()

    # 9 Testing
    page(doc, 9, "Testing Deployment and Operations", "Demonstrate that the product works and can be operated reliably.")
    doc.add_heading("Test Strategy", level=2)
    table(doc, ["TEST TYPE", "SCOPE", "TOOL OR METHOD", "RESULT OR EVIDENCE"], [
        ["Syntax", "Python routes and services; browser JavaScript", "py_compile and node --check", "Passes for current project changes"],
        ["Public API", "Health, advisory and market route contracts", "Pytest public-surface suite when installed", "Project includes public surface tests"],
        ["Live integration", "Brinjal Odisha mandi lookup", "AGMARKNET 2.0 call and UI verification", "Attabira APMC quote returned on 14 Sep 2026"],
        ["Visual QA", "Farmer journey and responsive UI", "Browser accessibility tree and UI review", "Crop, state and result visible in corrected demo"],
    ], [1.1, 1.9, 1.95, 1.65])
    doc.add_heading("Critical Test Cases", level=2)
    table(doc, ["ID", "SCENARIO", "EXPECTED RESULT", "STATUS"], [
        ["TC-01", "Brinjal with Odisha", "Official mandi range, date and source displayed", "Pass"],
        ["TC-02", "No market record", "Clear unavailable message; no guessed price", "Pass"],
        ["TC-03", "Missing leaf photo", "Form blocks submission and explains required input", "Pass"],
        ["TC-04", "Location advisory", "Weather profile fills only after user permission", "Pass"],
        ["TC-05", "Other crop advisory", "Custom crop field appears and is validated", "Pass"],
    ], [0.7, 2.1, 2.65, 1.2])
    doc.add_page_break()

    # 10 Impact
    page(doc, 10, "Impact Roadmap and Final Handover", "Close the document as a product team would with evidence ownership and next steps.")
    doc.add_heading("Results and Impact", level=2)
    doc.add_paragraph("The prototype demonstrates an end-to-end farmer assistance flow across disease screening, verified mandi-price lookup, location-aware advisory and WhatsApp support. A live Odisha Brinjal lookup was verified against AGMARKNET 2.0 on 14 September 2026: Attabira APMC, ₹2,650 minimum, ₹2,800 modal and ₹2,950 maximum per quintal.")
    table(doc, ["OUTCOME OR KPI", "RESULT", "EVIDENCE"], [
        ["Verified market result", "Brinjal price returned for Odisha", "AGMARKNET 2.0 result with mandi and publication date"],
        ["Farmer-friendly guidance", "Hindi and English interaction available", "Language selector, translated labels and browser speech support"],
        ["Consent-aware design", "Optional declaration and pilot consent", "Visible consent card; core features remain usable"],
    ], [1.8, 2.25, 2.6])
    doc.add_heading("Risks Limitations and Roadmap", level=2)
    table(doc, ["RISK OR LIMITATION", "IMPACT", "MITIGATION OR NEXT STEP"], [
        ["Government feed latency or absence", "Price may be unavailable", "Show transparent status, recent published date and retry guidance"],
        ["Location does not provide field-specific soil lab results", "NPK and pH are estimates", "Use regional baseline only; direct farmer to recent soil test"],
        ["Pilot data scale", "Limited aggregate insights initially", "Keep consent explicit and expand through responsible pilot onboarding"],
    ], [1.9, 1.7, 3.05])
    table(doc, ["ROADMAP PHASE", "TIMEFRAME", "KEY ENHANCEMENT", "EXPECTED IMPACT"], [
        ["Now", "Hackathon", "Live market verification and farmer-first web flow", "Reliable demo of core value"],
        ["Next", "Pilot", "Broader states, local crop aliases and advisor feedback", "Higher regional relevance"],
        ["Future", "Scale", "Production hosting, authenticated partners and field soil-report import", "Personalised trusted guidance"],
    ], [1.1, 1.1, 2.55, 1.9])
    doc.add_page_break()

    # 11 Handover
    page(doc, 11, "Handover and Submission Checklist", "Final product and engineering deliverables for the hackathon jury.")
    table(doc, ["ITEM", "STATUS OR LINK"], [
        ["Source repository", "KrishiMitr local repository"],
        ["Live demo", "Local demo with FastAPI server and responsive web interface"],
        ["WhatsApp integration", "Twilio webhook route included; credentials configured externally"],
        ["Architecture record", "Included in this document"],
        ["Mandi-price evidence", "AGMARKNET 2.0 verified Brinjal Odisha lookup"],
        ["Testing record", "Syntax, API and UI checks described in Section 09"],
        ["References and acknowledgements", "AGMARKNET 2.0, Open-Meteo, Twilio, PlantVillage-compatible disease model assets"],
        ["Mentor approval", "To be completed by mentor"],
    ], [2.25, 4.25])
    doc.add_paragraph()
    p = doc.add_paragraph("Mentor Sign-off: ______________________________    Date: ______________")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    main()
